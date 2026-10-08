"""Standardize an existing Word document with a Word template."""

from __future__ import annotations

import argparse
import json
import re
import sys
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn as docx_qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from build_docx import (
    CT_NS,
    ET,
    REL_NS,
    ConversionError,
    ConfirmationRequired,
    analyze_template,
    build_document_xml,
    classify_prefix,
    count_elements,
    update_styles_numbering,
    write_docx,
)


HEADING_NAME = re.compile(r"(?:heading|标题)\s*(\d+)", re.IGNORECASE)
WORD_NS = docx_qn("w:p").split("}", 1)[0][1:]


def paragraph_text(paragraph: Paragraph) -> str:
    parts: list[str] = []
    for node in paragraph._p.iter():
        local_name = node.tag.rsplit("}", 1)[-1]
        if local_name == "t":
            parts.append(node.text or "")
        elif local_name == "br":
            parts.append("\n")
        else:
            parts.append("\t")
    return "".join(parts)


def style_name(paragraph: Paragraph) -> str:
    try:
        return paragraph.style.name or ""
    except Exception:
        return ""


def heading_level(paragraph: Paragraph) -> int | None:
    name = style_name(paragraph).lower()
    if name in {"title", "标题", "文档标题"}:
        return 1
    match = HEADING_NAME.search(style_name(paragraph))
    if match:
        return min(4, max(1, int(match.group(1))))
    ppr = paragraph._p.pPr
    if ppr is None:
        return None
    outline = ppr.find(docx_qn("w:outlineLvl"))
    if outline is None:
        return None
    value = outline.get(docx_qn("w:val"))
    if value is None or not value.isdigit():
        return None
    return min(4, max(1, int(value) + 1))


def heading_title_hint(paragraph: Paragraph) -> bool:
    if style_name(paragraph).lower() in {"title", "标题", "文档标题"}:
        return True
    ppr = paragraph._p.pPr
    if ppr is None:
        return False
    numpr = ppr.find(docx_qn("w:numPr"))
    if numpr is None:
        return False
    numid = numpr.find(docx_qn("w:numId"))
    return numid is not None and numid.get(docx_qn("w:val")) == "0"


def list_kind(paragraph: Paragraph) -> str | None:
    name = style_name(paragraph).lower()
    if any(token in name for token in ("list bullet", "bullet list", "项目符号")):
        return "unordered"
    if any(token in name for token in ("list number", "numbered list", "列表编号")):
        return "ordered"
    ppr = paragraph._p.pPr
    if ppr is not None and ppr.find(docx_qn("w:numPr")) is not None:
        return "ordered"
    text = paragraph_text(paragraph).strip()
    if re.match(r"^[-*+]\s+", text):
        return "unordered"
    if re.match(r"^\d+[.)、]\s+", text):
        return "ordered"
    return None


def strip_list_prefix(text: str, kind: str) -> str:
    if kind == "unordered":
        return re.sub(r"^[-*+]\s+", "", text, count=1)
    return re.sub(r"^\d+[.)、]\s+", "", text, count=1)


def is_quote(paragraph: Paragraph) -> bool:
    name = style_name(paragraph).lower()
    return "quote" in name or "引用" in style_name(paragraph)


def is_code(paragraph: Paragraph) -> bool:
    name = style_name(paragraph).lower()
    return "code" in name or "preformatted" in name or "代码" in style_name(paragraph)


def extracted_images(paragraph: Paragraph, document: Any, image_dir: Path) -> list[tuple[str, Path]]:
    embed_attribute = docx_qn("r:embed")
    blip_tag = docx_qn("a:blip")
    relationships = [
        node.get(embed_attribute)
        for node in paragraph._p.iter()
        if node.tag == blip_tag and node.get(embed_attribute)
    ]
    images: list[tuple[str, Path]] = []
    plain_text = paragraph_text(paragraph).strip()
    fallback_alt = plain_text.splitlines()[0] if plain_text else ""
    for relationship_id in relationships:
        try:
            part = document.part.related_parts[relationship_id]
            payload = part.blob
            suffix = Path(str(part.partname)).suffix or ".png"
        except (KeyError, AttributeError):
            continue
        path = image_dir / f"extracted-{len(images) + 1}{suffix}"
        path.write_bytes(payload)
        alt = fallback_alt or f"图片 {len(images) + 1}"
        images.append((alt, path))
    return images


def table_block(table: Table) -> dict[str, Any]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            cell_text = "<br>".join(
                paragraph_text(paragraph).replace("\n", "<br>")
                for paragraph in cell.paragraphs
            )
            cells.append(cell_text.strip("<br>").strip())
        rows.append(cells)
    width = max((len(row) for row in rows), default=1)
    rows = [row + [""] * (width - len(row)) for row in rows]
    has_header = len(rows) > 1 and any(value.strip() for value in rows[0])
    return {"type": "table", "rows": rows, "has_header": has_header}


def extract_blocks(input_path: Path, image_dir: Path) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        document = Document(str(input_path))
    except Exception as error:
        raise ConversionError(f"Word 文件无法打开：{input_path}（{error}）") from error

    blocks: list[dict[str, Any]] = []
    warnings: list[str] = []
    pending_list: dict[str, Any] | None = None

    def flush_list() -> None:
        nonlocal pending_list
        if pending_list:
            blocks.append(pending_list)
            pending_list = None

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == docx_qn("w:p"):
            paragraph = Paragraph(child, document)
            text = paragraph_text(paragraph).strip()
            images = extracted_images(paragraph, document, image_dir)
            if not text and not images:
                continue
            if images:
                flush_list()
                if text:
                    blocks.append({"type": "paragraph", "text": text.replace("\n", "<br>")})
                for alt, path in images:
                    blocks.append({"type": "image", "alt": alt, "src": str(path)})
                continue

            level = heading_level(paragraph)
            if level is not None:
                flush_list()
                title_hint = heading_title_hint(paragraph)
                title = text.replace("\n", " ")
                number_hint = None
                classification = classify_prefix(title)
                if classification:
                    _, prefix, title = classification
                    number_hint = prefix
                blocks.append({
                    "type": "heading",
                    "level": level,
                    "source_level": level,
                    "text": text,
                    "title_hint": title_hint,
                    "title": title,
                    "number_hint": number_hint,
                })
                continue

            kind = list_kind(paragraph)
            if kind:
                item_text = strip_list_prefix(text, kind).replace("\n", "<br>")
                if pending_list and pending_list["ordered"] == (kind == "ordered"):
                    pending_list["items"].append({"text": item_text, "marker": None})
                else:
                    flush_list()
                    pending_list = {
                        "type": "list",
                        "ordered": kind == "ordered",
                        "items": [{"text": item_text, "marker": None}],
                    }
                continue

            flush_list()
            if is_code(paragraph):
                blocks.append({"type": "code", "text": text})
            elif is_quote(paragraph):
                blocks.append({"type": "quote", "text": text.replace("\n", "<br>")})
            else:
                blocks.append({"type": "paragraph", "text": text.replace("\n", "<br>")})
        elif child.tag == docx_qn("w:tbl"):
            flush_list()
            blocks.append(table_block(Table(child, document)))
        elif child.tag not in {docx_qn("w:sectPr")}:
            warnings.append("文档包含暂不识别的内容块，已保留到诊断信息中。")

    flush_list()
    headings = [block for block in blocks if block["type"] == "heading"]
    if (
        len(headings) > 1
        and headings[0].get("title_hint")
        and len({block.get("source_level") for block in headings}) == 1
    ):
        headings[0]["level"] = 1
        for block in headings[1:]:
            block["level"] = min(4, headings[0]["source_level"] + 1)
    if not blocks:
        raise ConversionError("Word 文档中没有可识别的正文内容。")
    return blocks, warnings


def check_heading_consistency(blocks: list[dict[str, Any]]) -> list[str]:
    signatures: dict[str, set[int]] = {}
    for block in blocks:
        if block["type"] != "heading" or not block.get("number_hint"):
            continue
        signature, prefix, _title = classify_prefix(block["text"])
        if not signature:
            continue
        if ":" in signature:
            signature = f"{signature.split(':', 1)[0]}:{prefix.count('.') + 1}"
        signatures.setdefault(signature, set()).add(block["level"])
    return [
        f"同类章节编号被用于不同标题层级：{signature}"
        for signature, levels in signatures.items()
        if len(levels) > 1
    ]


def document_counts(path: Path) -> dict[str, int]:
    document = Document(str(path))
    paragraphs = [
        Paragraph(element, document)
        for element in document.element.body.iterchildren(docx_qn("w:p"))
    ]
    tables = list(document.element.body.iterchildren(docx_qn("w:tbl")))
    image_tag = docx_qn("a:blip")
    images = [node for node in document.element.body.iter() if node.tag == image_tag]
    return {
        "paragraphs": sum(1 for paragraph in paragraphs if paragraph_text(paragraph).strip()),
        "tables": len(tables),
        "images": len(images),
    }


def standardize(
    input_path: Path,
    output_path: Path,
    template_path: Path,
    level: str,
) -> dict[str, Any]:
    with tempfile.TemporaryDirectory(prefix="md-to-word-") as temp_dir:
        blocks, extraction_warnings = extract_blocks(input_path, Path(temp_dir))
        heading_levels = {block["level"] for block in blocks if block["type"] == "heading"}
        template = analyze_template(template_path, heading_levels)
        update_styles_numbering(
            template["trees"]["word/styles.xml"],
            template["headings"],
            template["heading_numbering"],
        )
        rels_root = template["trees"].get(
            "word/_rels/document.xml.rels",
            ET.Element(REL_NS + "Relationships"),
        )
        content_types_root = template["trees"].get(
            "[Content_Types].xml",
            ET.Element(CT_NS + "Types"),
        )
        media: list[tuple[str, bytes, str]] = []
        document_root = build_document_xml(
            blocks,
            template,
            input_path.parent,
            media,
            rels_root,
            content_types_root,
        )
        write_docx(
            output_path,
            template_path,
            document_root,
            template["trees"]["word/styles.xml"],
            template["trees"].get("word/numbering.xml"),
            rels_root,
            content_types_root,
            media,
        )

    before = document_counts(input_path)
    after = document_counts(output_path)
    warnings = list(extraction_warnings) + check_heading_consistency(blocks)
    if before["tables"] != after["tables"]:
        warnings.append(f"表格数量变化：原文 {before['tables']}，输出 {after['tables']}。")
    if before["images"] != after["images"]:
        warnings.append(f"图片数量变化：原文 {before['images']}，输出 {after['images']}。")
    if before["paragraphs"] > after["paragraphs"]:
        warnings.append(
            f"非空段落数减少：原文 {before['paragraphs']}，输出 {after['paragraphs']}，请人工确认。"
        )

    return {
        "input": str(input_path),
        "output": str(output_path),
        "template": str(template_path),
        "level": level,
        "intermediate_markdown": None,
        "counts": count_elements(document_root),
        "content_counts": {"source": before, "output": after},
        "numbering_source": {
            str(heading_level_key): info["source"]
            for heading_level_key, info in template["heading_numbering"].items()
        },
        "warnings": warnings,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Word input path")
    parser.add_argument("output", type=Path, nargs="?", help="Standardized DOCX output path")
    parser.add_argument("--template", type=Path, help="Word template path")
    parser.add_argument(
        "--level",
        choices=("auto", "style", "structure", "migrate"),
        default="auto",
        help="Word processing level",
    )
    parser.add_argument("--overwrite", action="store_true", help="Overwrite the output DOCX")
    parser.add_argument("--json", action="store_true", help="Print a JSON result report")
    args = parser.parse_args()

    skill_dir = Path(__file__).resolve().parent.parent
    input_path = args.input.resolve()
    if not input_path.is_file() or input_path.suffix.lower() != ".docx":
        raise ConversionError(f"Word 文件不存在或不是 .docx：{input_path}")
    output_path = (
        args.output or input_path.with_name(f"{input_path.stem}.standardized.docx")
    ).resolve()
    if output_path.exists() and not args.overwrite:
        raise ConfirmationRequired(f"输出文件已存在：{output_path}。请确认是否覆盖。")
    template_path = (args.template or skill_dir / "templates" / "default.docx").resolve()
    level = "style" if args.level == "auto" else args.level
    result = standardize(input_path, output_path, template_path, level)

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print("Word 结构已识别并按模板标准化。")
        print(f"已使用 Word 模板：{template_path}")
        print("Word 样式与自动编号已应用。")
        print(f"Word 文档已生成并完成基础检查：{output_path}")
        if result["warnings"]:
            print("需要确认：")
            for warning in result["warnings"]:
                print(f"- {warning}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except ConfirmationRequired as error:
        print(f"需要用户确认：{error}", file=sys.stderr)
        raise SystemExit(2)
    except ConversionError as error:
        print(f"标准化失败：{error}", file=sys.stderr)
        raise SystemExit(1)
